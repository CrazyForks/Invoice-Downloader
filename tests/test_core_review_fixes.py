from pathlib import Path
from datetime import datetime, timedelta, timezone
import json
import shutil
import socket
import ssl
import threading
import time
import zipfile
from types import SimpleNamespace

import pytest
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import ec
from cryptography.x509.oid import NameOID
from openpyxl import load_workbook

import app_api
import email_fetcher
from app_api import InvoiceAppAPI, build_business_record_key, record_business_success
from app_archive_adapter import AppArchiveAdapter
from archive_service import ArchiveService
from candidate_pipeline import CandidatePipeline, CandidatePreflight
from email_fetcher import EmailFetcher
from extraction_pipeline import ExtractionOutcome
from invoice_extractor import InvoiceExtractor
from mailbox_scanner import MailboxScanError, MailboxScanner


def test_excel_export_keeps_external_text_literal_and_numeric_totals(tmp_path):
    api = InvoiceAppAPI()
    api.processed_invoices = [{
        "date": "=1+1", "amount": "+1+1", "merchant": "=1+1",
        "category": "=1+1", "path": "=1+1",
    }, {"date": "中文", "amount": "¥ 100.00", "merchant": "name\x01", "category": "正常", "path": "C:/票据.pdf"},
        {"date": "-1+1", "amount": "@SUM(1)", "merchant": "+1+1", "category": "-1+1", "path": "@file"}]
    api.error_invoices = [{
        "status": "=1+1", "reason": "=1+1", "date": "=1+1",
        "amount": "=1+1", "merchant": "=1+1", "name": "=1+1", "path": "=1+1",
    }]

    result = api.export_run_summary(str(tmp_path))
    assert result["success"]
    workbook = load_workbook(result["path"], data_only=False)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
    assert workbook["成功明细"]["C2"].value == "=1+1"
    assert workbook["成功明细"]["C3"].value == "name\ufffd"
    for index, expected in enumerate(("-1+1", "@SUM(1)", "+1+1", "-1+1", "@file"), 1):
        cell = workbook["成功明细"].cell(4, index)
        assert cell.value == expected and cell.data_type == "s"
    summary = {row[0].value: (row[1].value, row[2].value) for row in list(workbook["分类汇总"].rows)[1:]}
    assert summary["正常"] == (1, 100)
    with zipfile.ZipFile(result["path"]) as archive:
        assert not any(b"<f>" in archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/"))


def test_imap_connections_verify_identity_and_have_timeout(monkeypatch):
    seen = []

    class Mail:
        def __init__(self, *args, **kwargs):
            seen.append((args, kwargs))

        def login(self, *_args):
            return "OK", []

        def logout(self):
            return "BYE", []

    monkeypatch.setattr(email_fetcher.imaplib, "IMAP4_SSL", Mail)
    fetcher = EmailFetcher("a@qq.com", "fake-secret")
    assert fetcher.connect()
    fetcher.disconnect()
    api = InvoiceAppAPI()
    assert api.test_email_auth("a@qq.com", "fake-secret")["success"]
    assert len(seen) == 2
    for _, kwargs in seen:
        context = kwargs["ssl_context"]
        assert context.verify_mode == ssl.CERT_REQUIRED
        assert context.check_hostname
        assert 0 < kwargs["timeout"] <= 60


def test_imap_certificate_failure_never_logs_in(monkeypatch):
    logins = []

    def rejected(*_args, **_kwargs):
        raise ssl.SSLCertVerificationError("certificate verify failed")

    monkeypatch.setattr(email_fetcher.imaplib, "IMAP4_SSL", rejected)
    fetcher = EmailFetcher("a@qq.com", "fake-secret")
    monkeypatch.setattr(fetcher, "_send_imap_id_command", lambda: logins.append("id"))
    assert not fetcher.connect()
    assert fetcher.mail is None
    assert not logins
    assert fetcher.certificate_error
    assert "证书验证失败" in InvoiceAppAPI().test_email_auth("a@qq.com", "fake-secret")["message"]


def test_imap_login_timeout_closes_socket_and_reports_timeout(monkeypatch, tmp_path):
    mails = []

    class Mail:
        def __init__(self, *_args, **_kwargs):
            self.closed = False
            mails.append(self)

        def login(self, *_args):
            raise TimeoutError("silent server")

        def shutdown(self):
            self.closed = True

    monkeypatch.setattr(email_fetcher.imaplib, "IMAP4_SSL", Mail)
    fetcher = EmailFetcher("a@qq.com", "fake-secret", staging_dir=str(tmp_path))
    assert not fetcher.connect()
    assert fetcher.timeout_error
    assert fetcher.mail is None
    assert mails[0].closed
    assert "超时" in InvoiceAppAPI().test_email_auth("a@qq.com", "fake-secret")["message"]
    assert all(mail.closed for mail in mails)


@pytest.mark.parametrize("certificate_case", ["trusted", "untrusted", "wrong-host", "expired"])
def test_local_imap_tls_verifies_certificate_before_login(tmp_path, monkeypatch, certificate_case):
    now = datetime.now(timezone.utc)

    def make_ca(name):
        key = ec.generate_private_key(ec.SECP256R1())
        subject = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, name)])
        cert = (x509.CertificateBuilder().subject_name(subject).issuer_name(subject)
                .public_key(key.public_key()).serial_number(x509.random_serial_number())
                .not_valid_before(now - timedelta(days=2)).not_valid_after(now + timedelta(days=30))
                .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
                .sign(key, hashes.SHA256()))
        return key, cert

    ca_key, ca_cert = make_ca("Test IMAP CA")
    leaf_key = ec.generate_private_key(ec.SECP256R1())
    leaf_name = "wrong.example" if certificate_case == "wrong-host" else "localhost"
    subject = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, leaf_name)])
    end = now - timedelta(days=1) if certificate_case == "expired" else now + timedelta(days=1)
    start = end - timedelta(days=2) if certificate_case == "expired" else now - timedelta(days=1)
    leaf_cert = (x509.CertificateBuilder().subject_name(subject).issuer_name(ca_cert.subject)
                 .public_key(leaf_key.public_key()).serial_number(x509.random_serial_number())
                 .not_valid_before(start).not_valid_after(end)
                 .add_extension(x509.SubjectAlternativeName([x509.DNSName(leaf_name)]), critical=False)
                 .sign(ca_key, hashes.SHA256()))
    cert_path, key_path, ca_path = (tmp_path / name for name in ("server.pem", "server.key", "ca.pem"))
    cert_path.write_bytes(leaf_cert.public_bytes(serialization.Encoding.PEM))
    key_path.write_bytes(leaf_key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()))
    trusted_ca = make_ca("Other CA")[1] if certificate_case == "untrusted" else ca_cert
    ca_path.write_bytes(trusted_ca.public_bytes(serialization.Encoding.PEM))

    server_context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    server_context.load_cert_chain(str(cert_path), str(key_path))
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    listener.listen(1)
    listener.settimeout(5)
    login_seen = []
    failures = []
    commands = []

    def serve():
        try:
            connection, _ = listener.accept()
            with connection:
                with server_context.wrap_socket(connection, server_side=True) as tls:
                    tls.settimeout(5)
                    tls.sendall(b"* OK local test IMAP ready\r\n")
                    stream = tls.makefile("rb")
                    while line := stream.readline():
                        commands.append(line)
                        tag, command, *_ = line.split(b" ", 2)
                        command = command.strip().upper()
                        if command == b"CAPABILITY":
                            tls.sendall(b"* CAPABILITY IMAP4rev1\r\n" + tag + b" OK capabilities\r\n")
                        elif command == b"LOGIN":
                            login_seen.append(True)
                            tls.sendall(tag + b" OK logged in\r\n")
                        elif command == b"LOGOUT":
                            tls.sendall(b"* BYE leaving\r\n" + tag + b" OK logout\r\n")
                            return
        except (ssl.SSLError, OSError) as exc:
            failures.append(type(exc).__name__)
        finally:
            listener.close()

    thread = threading.Thread(target=serve)
    thread.start()
    original_context = ssl.create_default_context
    monkeypatch.setattr(email_fetcher.ssl, "create_default_context", lambda: original_context(cafile=str(ca_path)))
    if certificate_case == "trusted":
        monkeypatch.setattr(email_fetcher, "IMAP_TIMEOUT_SECONDS", 1)
    fetcher = EmailFetcher("a@qq.com", "fake-secret", imap_server="localhost", imap_port=listener.getsockname()[1], staging_dir=str(tmp_path / "staging"))
    try:
        assert fetcher.connect() is (certificate_case == "trusted"), (commands, failures)
        assert bool(login_seen) is (certificate_case == "trusted")
        if certificate_case == "trusted":
            started = time.monotonic()
            with pytest.raises(MailboxScanError):
                fetcher.fetch_emails_by_date("2026-06-01", "2026-06-02")
            assert time.monotonic() - started < 3
            assert fetcher.mail is None
    finally:
        fetcher.disconnect()
        thread.join(timeout=6)
    assert not thread.is_alive()


def test_163_id_eof_stops_after_first_read():
    class Mail:
        def __init__(self):
            self.reads = 0

        def _new_tag(self):
            return b"A001"

        def send(self, *_args):
            pass

        def readline(self):
            self.reads += 1
            if self.reads > 1:
                raise AssertionError("read after EOF")
            return b""

    fetcher = EmailFetcher("a@163.com", "fake-secret")
    fetcher.mail = Mail()
    with pytest.raises(ConnectionError):
        fetcher._send_imap_id_command()
    assert fetcher.mail.reads == 1


def test_imap_abort_wakes_blocked_socket_read(tmp_path):
    reader, peer = socket.socketpair()
    fetcher = EmailFetcher("a@qq.com", "fake-secret", staging_dir=str(tmp_path))
    fetcher.mail = type("Mail", (), {"sock": reader})()
    entered = threading.Event()
    completed = threading.Event()

    def blocked_read():
        entered.set()
        try:
            reader.recv(1)
        except OSError:
            pass
        completed.set()

    thread = threading.Thread(target=blocked_read)
    thread.start()
    assert entered.wait(2)
    try:
        fetcher.abort()
        assert completed.wait(2)
    finally:
        peer.close()
        thread.join(timeout=2)
    assert not thread.is_alive()


@pytest.mark.parametrize("stage", ["select", "search", "fetch"])
def test_mailbox_socket_timeout_never_reports_success(stage):
    class Mail:
        def select(self, *_args, **_kwargs):
            if stage == "select":
                raise TimeoutError("silent SELECT")
            return "OK", []

        def uid(self, command, *_args):
            if command == "SEARCH":
                if stage == "search":
                    raise TimeoutError("silent SEARCH")
                return "OK", [b"123"]
            raise TimeoutError("silent FETCH")

    scanner = MailboxScanner(Mail())
    with pytest.raises(MailboxScanError):
        scanner.scan(datetime(2026, 6, 1).date(), datetime(2026, 6, 2).date())


def test_downloaded_url_manual_review_preserves_pdf_before_staging_cleanup(tmp_path):
    api = InvoiceAppAPI()
    staging = tmp_path / "staging"
    staging.mkdir()
    source = staging / "invoice.pdf"
    source.write_bytes(b"%PDF-1.4\nsynthetic")
    api._active_staging_path = lambda: str(staging)
    result = api._send_to_manual_check(
        str(tmp_path / "output"), str(source), "COMPANY_PURCHASER_UNKNOWN",
        metadata={"source_kind": "url", "file_name": "invoice.pdf"}, is_url=True,
    )
    source.unlink()
    assert Path(result).read_bytes() == b"%PDF-1.4\nsynthetic"
    assert Path(result).suffix == ".pdf"


def test_downloaded_url_passes_preflight_archive_and_cleanup_with_pdf_intact(tmp_path):
    staging = tmp_path / "staging"
    staging.mkdir()
    downloaded = staging / "downloaded.pdf"
    downloaded.write_bytes(b"%PDF-1.4\nsynthetic receipt")
    api = InvoiceAppAPI()
    api._active_staging_path = lambda: str(staging)
    candidate = CandidatePipeline().collect([{
        "filepath": "https://provider.example/invoice?token=secret",
        "email_id": "mail-1", "source_kind": "url",
    }])[0]

    class Converter:
        def process_invoice_links(self, *_args, **_kwargs):
            return [{"status": "success", "pdf_path": str(downloaded)}]

    class Extractor:
        def probe_local_only(self, path, **_kwargs):
            assert path == str(downloaded)
            return SimpleNamespace(status="resolved", result={"Type": "发票", "Purchaser": ""}, engine="local", reason_code="LOCAL")

    preflight = CandidatePreflight(
        api=api, extractor=Extractor(), working_history=set(), sidecar={},
        sidecar_lock=threading.Lock(), converter_factory=Converter,
    )
    outcome = preflight(candidate)
    assert outcome.status == "resolved"
    assert outcome.candidate.identity.source_kind == "url"
    assert outcome.payload["pdf_path"] == str(downloaded)

    output = tmp_path / "output"
    adapter = AppArchiveAdapter(api=api, extractor=object(), save_path=str(output), business_records={}, trace_store=object())
    service = ArchiveService(
        normalizer=lambda value: {**value.to_legacy_payload(), "archive_status": "manual_review", "manual_reason_code": "COMPANY_PURCHASER_UNKNOWN"},
        classifier=adapter.classify, archive_operation=adapter.archive_operation,
        dedupe_key=adapter.dedupe_key,
    )
    report = service.archive([outcome], output)
    assert report.manual_count == 1
    retained = Path(report.outcomes[0].archive_path)
    assert retained.read_bytes() == downloaded.read_bytes()
    api._cleanup_temp_folders(staging_dir=staging)
    assert retained.read_bytes() == b"%PDF-1.4\nsynthetic receipt"
    assert not staging.exists()


def test_url_manual_review_uses_downloaded_artifact_path(tmp_path):
    api = InvoiceAppAPI()
    staging = tmp_path / "staging"
    staging.mkdir()
    source = staging / "downloaded.pdf"
    source.write_bytes(b"%PDF-1.4\nfrom URL")
    api._active_staging_path = lambda: str(staging)
    candidate = CandidatePipeline().collect([{
        "filepath": "https://example.test/invoice?token=secret",
        "source_url": "https://example.test/invoice?token=secret",
        "email_id": "mail-1",
    }])[0]
    outcome = ExtractionOutcome(
        candidate=candidate, status="manual_review",
        reason_code="EXTRACTOR_ALL_ENGINES_FAILED", artifact_path=str(source),
    )
    adapter = AppArchiveAdapter(
        api=api, extractor=object(), save_path=str(tmp_path / "output"),
        business_records={}, trace_store=object(),
    )
    decision = adapter.archive_operation(outcome, None, "", tmp_path / "output")
    source.unlink()
    assert Path(decision.path).read_bytes() == b"%PDF-1.4\nfrom URL"
    assert "secret" not in (tmp_path / "output").joinpath("待人工复核", Path(decision.path).name + ".json").read_text(encoding="utf-8")


def test_downloaded_url_manual_copy_failure_cannot_complete(tmp_path, monkeypatch):
    api = InvoiceAppAPI()
    staging = tmp_path / "staging"
    staging.mkdir()
    source = staging / "invoice.pdf"
    source.write_bytes(b"%PDF-1.4\n")
    api._active_staging_path = lambda: str(staging)
    candidate = CandidatePipeline().collect([{"filepath": "https://example.test/invoice", "email_id": "mail-1"}])[0]
    outcome = ExtractionOutcome(candidate, "manual_review", "NEEDS_REVIEW", artifact_path=str(source))
    output = tmp_path / "output"
    adapter = AppArchiveAdapter(api=api, extractor=object(), save_path=str(output), business_records={}, trace_store=object())
    monkeypatch.setattr(shutil, "copy2", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("copy failed")))
    report = ArchiveService(archive_operation=adapter.archive_operation).archive([outcome], output)
    assert not report.can_complete
    assert report.unresolved_count == 1
    assert not list((output / "待人工复核").glob("*.json"))


def test_manual_copy_failure_keeps_downloaded_original_after_desktop_finalizer(tmp_path, monkeypatch):
    from run_coordinator import RunCoordinator, RunRequest
    from run_lifecycle import RunState

    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    staging = tmp_path / "staging"
    handle = api._run_lifecycle.begin("manual-copy-failure", staging)
    source = staging / "invoice.pdf"
    original = b"%PDF-1.4\nrecoverable original"
    source.write_bytes(original)
    api._active_run_handle = handle
    api._active_temp_dir = tmp_path / "temp"
    api._run_state_store.reset(handle.run_id)
    request = RunRequest(handle.run_id, "2026-06-01", "2026-06-13", str(output), "", "account", "qq", run_root=str(tmp_path))
    candidate = CandidatePipeline().collect([
        {"filepath": "https://example.test/invoice", "email_id": "mail-1"}
    ])[0]
    outcome = ExtractionOutcome(candidate, "manual_review", "NEEDS_REVIEW", artifact_path=str(source))
    adapter = AppArchiveAdapter(api=api, extractor=object(), save_path=str(output), business_records={}, trace_store=object())
    deps = api._build_run_dependencies(request, email_address="a@qq.com", auth_code="x", api_key="y")
    fetcher = SimpleNamespace(disconnect=lambda: None)
    deps.connect = lambda _request: fetcher
    deps.scan = lambda *_args: ["mail"]
    deps.candidate = lambda *_args: [candidate]
    deps.extract = lambda *_args: [outcome]
    deps.archive = lambda outcomes, _request: ArchiveService(
        archive_operation=adapter.archive_operation
    ).archive(outcomes, output)
    monkeypatch.setattr(shutil, "copy2", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("copy failed")))

    result = RunCoordinator(api._run_lifecycle, api._run_state_store, deps).run(request, handle=handle)
    assert result.state is RunState.FAILED
    assert result.reason_code == "ARCHIVE_INCOMPLETE"
    assert source.read_bytes() == original
    assert any(str(source) in entry.get("msg", "") for entry in api.logs)
    assert not list((output / "待人工复核").glob("*.json"))


def test_business_records_restore_from_output_scoped_state(tmp_path, monkeypatch):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    state = tmp_path / "state"
    monkeypatch.setattr(api, "_output_state_dir", lambda _path: str(state))
    records = {"_26000000000000000001::invoice": {"file": "invoice.pdf", "artifact_role": "invoice"}}
    api._commit_output_state(str(state), {"history-key"}, records)
    extractor = InvoiceExtractor(api_key="", output_dir=str(output))
    session = api._create_processing_pipeline_session([], "", str(output), _extractor=extractor)
    try:
        assert session._business_records == records
    finally:
        session.close()


@pytest.mark.parametrize("interruption", ["failed", "cancelled", "commit_error"])
def test_completed_old_state_survives_first_failed_upgraded_run(tmp_path, monkeypatch, interruption):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    state = tmp_path / "state"
    state.mkdir()
    monkeypatch.setattr(api, "_output_state_dir", lambda _path: str(state))
    records = {"old-invoice": {"file": "invoice.pdf"}}
    (state / "processed_records.json").write_text(json.dumps(records), encoding="utf-8")
    (state / ".antigravity_history.json").write_text(json.dumps(["old-history"]), encoding="utf-8")
    (state / "run_state.json").write_text(json.dumps({"status": "completed"}), encoding="utf-8")

    extractor = InvoiceExtractor(api_key="", output_dir=str(output))
    first = api._create_processing_pipeline_session([], "", str(output), _extractor=extractor)
    try:
        assert first._business_records == records
        assert "old-history" in first._working_history
        if interruption == "commit_error":
            real_write = api._write_json_file

            def fail_bundle(path, payload):
                if path.endswith("committed_state.json"):
                    raise OSError("synthetic commit failure")
                return real_write(path, payload)

            monkeypatch.setattr(api, "_write_json_file", fail_bundle)
            with pytest.raises(OSError, match="synthetic commit failure"):
                api._commit_output_state(str(state), {"new-history"}, {"new-invoice": {"file": "new.pdf"}})
        else:
            api._mark_output_run_state(str(state), interruption)
    finally:
        first.close()

    assert (state / "committed_state.json").exists()
    restarted = InvoiceAppAPI()
    monkeypatch.setattr(restarted, "_output_state_dir", lambda _path: str(state))
    second = restarted._create_processing_pipeline_session(
        [], "", str(output), _extractor=InvoiceExtractor(api_key="", output_dir=str(output))
    )
    try:
        assert second._business_records == records
        assert "old-history" in second._working_history
    finally:
        second.close()


@pytest.mark.parametrize("status", ["failed", "cancelled"])
def test_incomplete_old_state_is_not_migrated(tmp_path, monkeypatch, status):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    state = tmp_path / "state"
    state.mkdir()
    monkeypatch.setattr(api, "_output_state_dir", lambda _path: str(state))
    (state / "processed_records.json").write_text(json.dumps({"uncommitted": {"file": "x.pdf"}}), encoding="utf-8")
    (state / ".antigravity_history.json").write_text(json.dumps(["uncommitted"]), encoding="utf-8")
    (state / "run_state.json").write_text(json.dumps({"status": status}), encoding="utf-8")
    session = api._create_processing_pipeline_session(
        [], "", str(output), _extractor=InvoiceExtractor(api_key="", output_dir=str(output))
    )
    try:
        assert session._business_records == {}
        assert "uncommitted" not in session._working_history
        assert not (state / "committed_state.json").exists()
    finally:
        session.close()


def test_distinct_bytes_of_same_invoice_dedupe_across_runs(tmp_path):
    output = tmp_path / "output"
    output.mkdir()
    api = InvoiceAppAPI()
    state = api._output_state_dir(str(output))
    sources = []
    for index in range(3):
        path = tmp_path / f"version-{index}.pdf"
        path.write_bytes(f"%PDF-1.4\nversion-{index}".encode())
        sources.append(path)
    candidates = CandidatePipeline().collect([{"filepath": str(path), "email_id": f"mail-{index}"} for index, path in enumerate(sources)])
    assert len({item.identity.document_id for item in candidates}) == 3
    written = []
    info = {"InvoiceNumber": "26000000000000000001", "Type": "发票", "is_invoice": True}

    def make_service(records):
        def write(outcome, root):
            target = root / outcome.candidate.source_filename
            shutil.copy2(outcome.candidate.source_path, target)
            written.append(str(target))
            record_business_success(records, "", outcome.payload["info_json"]["InvoiceNumber"], outcome.payload["info_json"], target.name)
            return str(target)

        return ArchiveService(
            writer=write, existing_dedupe_keys=records,
            dedupe_key=lambda outcome, payload: build_business_record_key("", payload["info_json"]["InvoiceNumber"], payload["info_json"], outcome.candidate.source_filename),
        )

    first_records = {}
    first = make_service(first_records).archive([ExtractionOutcome.resolved(candidates[0], {"info_json": info})], output)
    assert first.archived_count == 1
    api._commit_output_state(state, set(), first_records)
    resumed = api._load_business_records(str(output), state)
    second = make_service(resumed).archive([ExtractionOutcome.resolved(candidates[1], {"info_json": info})], output)
    assert second.duplicate_count == 1
    distinct = {**info, "InvoiceNumber": "26000000000000000002"}
    third = make_service(resumed).archive([ExtractionOutcome.resolved(candidates[2], {"info_json": distinct})], output)
    assert third.archived_count == 1
    assert len(written) == 2

    other_output = tmp_path / "other-output"
    other_output.mkdir()
    assert api._load_business_records(str(other_output), api._output_state_dir(str(other_output))) == {}


def test_business_records_merge_legacy_and_scoped_history_without_overwriting(tmp_path, monkeypatch):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    state = tmp_path / "state"
    monkeypatch.setattr(api, "_output_state_dir", lambda _path: str(state))
    legacy = {"legacy-key": {"file": "old.pdf"}, "shared-key": {"file": "old-name.pdf"}}
    (output / "processed_records.json").write_text(json.dumps(legacy), encoding="utf-8")
    api._commit_output_state(str(state), {"history-key"}, {"shared-key": {"file": "new-name.pdf"}, "state-key": {"file": "new.pdf"}})
    extractor = InvoiceExtractor(api_key="", output_dir=str(output))
    session = api._create_processing_pipeline_session([], "", str(output), _extractor=extractor)
    try:
        assert session._business_records == {
            "legacy-key": {"file": "old.pdf"}, "shared-key": {"file": "new-name.pdf"},
            "state-key": {"file": "new.pdf"},
        }
        assert json.loads((state / "processed_records.legacy-backup.json").read_text(encoding="utf-8")) == legacy
    finally:
        session.close()


def test_corrupt_legacy_business_records_fail_closed(tmp_path):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    (output / "processed_records.json").write_text("{truncated", encoding="utf-8")
    with pytest.raises(json.JSONDecodeError):
        api._load_business_records(str(output), str(tmp_path / "state"))


def test_failed_state_commit_keeps_last_complete_business_snapshot(tmp_path, monkeypatch):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    output.mkdir()
    state = tmp_path / "state"
    monkeypatch.setattr(api, "_output_state_dir", lambda _path: str(state))
    api._commit_output_state(str(state), {"old-history"}, {"old-key": {"file": "old.pdf"}})
    real_write = api._write_json_file

    def fail_bundle(path, payload):
        if path.endswith("committed_state.json"):
            raise OSError("synthetic write failure")
        return real_write(path, payload)

    monkeypatch.setattr(api, "_write_json_file", fail_bundle)
    with pytest.raises(OSError, match="synthetic write failure"):
        api._commit_output_state(str(state), {"new-history"}, {"new-key": {"file": "new.pdf"}})
    extractor = InvoiceExtractor(api_key="", output_dir=str(output))
    session = api._create_processing_pipeline_session([], "", str(output), _extractor=extractor)
    try:
        assert session._business_records == {"old-key": {"file": "old.pdf"}}
        assert "old-history" in api._load_committed_history(str(state))
        assert "new-history" not in api._load_committed_history(str(state))
    finally:
        session.close()


@pytest.mark.parametrize(
    ("folder", "invoice_role", "companion_role"),
    [("打车", "ride_invoice", "ride_itinerary"), ("住宿发票", "hotel_invoice", "hotel_folio")],
)
def test_pair_rename_updates_results_export_and_business_filename(
    tmp_path, monkeypatch, folder, invoice_role, companion_role
):
    api = InvoiceAppAPI()
    output = tmp_path / "output"
    target = output / folder
    target.mkdir(parents=True)
    monkeypatch.setattr(api, "_output_state_dir", lambda _path: str(tmp_path / "state"))

    class Extractor:
        last_route_trace = {}

        def load_processed_records(self):
            return {}

        def route_and_rename_file(self, path, _info, custom_rules=None):
            return True, path

    session = api._create_processing_pipeline_session([], "", str(output), _extractor=Extractor())
    adapter = session._archive_service._finalizer.__self__
    trace = session._trace_store
    originals = []
    ids = []
    for index, role in enumerate((invoice_role, companion_role)):
        source = target / f"opaque-{index}.pdf"
        source.write_bytes(f"document-{index}".encode())
        candidate = CandidatePipeline().collect([{"filepath": str(source), "email_id": "same-mail"}])[0]
        document_id = candidate.identity.document_id
        ids.append(document_id)
        trace.start_document(source_filename=source.name, source_path=str(source), document_id=document_id)
        info = {
            "Date": "20260610", "Amount": "100.00", "Seller": "同一商户",
            "Purchaser": "辉瑞", "Type": "打车" if folder == "打车" else "住宿发票",
            "InvoiceNumber": f"2600000000000000000{index}", "is_invoice": True,
        }
        outcome = ExtractionOutcome.resolved(candidate, {"pdf_path": str(source), "metadata": candidate.to_legacy(), "info_json": info})
        decision = adapter.archive_operation(outcome, outcome.to_legacy_payload(), folder, output)
        assert decision.status == "archived"
        adapter.pairing_metadata[str(source)].update({
            "artifact_role": role, "source_message_uid": "same-mail", "date": "20260610",
            "amount": "100.00", "pairing_required": True,
        })
        originals.append(str(source))

    report = session.archive([])
    try:
        assert report.can_complete
        final_paths = {row["document_id"]: row["archive_target"] for row in trace.iter_records()}
        assert all(Path(path).is_file() for path in final_paths.values())
        assert any(final_paths[doc_id] != old for doc_id, old in zip(ids, originals))
        for row in api.get_results()["successInvoices"]:
            assert row["document_id"] in ids
            assert row["path"] == final_paths[row["document_id"]]
        exported = api.export_run_summary(str(tmp_path))["path"]
        workbook = load_workbook(exported)
        assert {workbook["成功明细"].cell(row, 5).value for row in (2, 3)} == set(final_paths.values())
        assert {record["file"] for record in session._business_records.values()} == {Path(path).name for path in final_paths.values()}
    finally:
        session.close()
